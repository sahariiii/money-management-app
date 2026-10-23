"""Write categorized rows to an .xlsx laid out like the budget workbook.

Columns A-F match the budget workbook's Expenses / Income columns B-G, so the
user can copy A2:F<n> and paste-as-values at B8 of their sheet. Review columns
come after F so they are easy to leave out of the copy.
"""

import io
from collections.abc import Sequence
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .workbook import BudgetSetup

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
REVIEW_FILL = PatternFill("solid", fgColor="F2F2F2")


def _as_date(value: str | None) -> date | str | None:
    try:
        return date.fromisoformat(value) if value else None
    except ValueError:
        return value  # keep whatever the model returned rather than dropping the row


def _write_sheet(ws, headers: list[str], rows: Sequence[Sequence], paste_cols: int, widths: list[int]) -> None:
    ws.append(headers)
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(1, col)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL if col <= paste_cols else REVIEW_FILL
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]
    for row in rows:
        ws.append(list(row))
    for (cell,) in ws.iter_rows(min_row=2, max_col=1):
        if isinstance(cell.value, date):
            cell.number_format = "mm-dd-yyyy"
    ws.freeze_panes = "A2"


def _category_dropdown(ws, column: str, list_range: str, last_row: int) -> None:
    if last_row < 2:
        return
    dv = DataValidation(type="list", formula1=list_range, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{column}2:{column}{last_row}")


def to_xlsx(
    expenses: Sequence[dict],
    income: Sequence[dict],
    skipped: Sequence[dict],
    setup: BudgetSetup,
) -> bytes:
    """Rows are dicts with the fields of the agent's output (plus 'statement' for the source file)."""
    wb = Workbook()

    lists = wb.create_sheet("Categories")
    for i, c in enumerate(setup.expense_categories, start=1):
        lists.cell(i, 1, c.name)
    for i, c in enumerate(setup.income_categories, start=1):
        lists.cell(i, 2, c.name)
    lists.sheet_state = "hidden"

    ws = wb.active
    ws.title = "Expenses"
    _write_sheet(
        ws,
        ["Date", "Store / Vendor", "$ Amount", "Expense Category", "Hide?", "Notes (Optional)",
         "Shared? (suggested)", "Statement description", "Review note", "Categorized by", "Statement"],
        [
            (_as_date(r["date"]), r["vendor"], r["amount"], r["category"], None, None,
             "yes" if r["likely_shared"] else None, r["statement_description"], r["review_note"],
             r.get("categorized_by"), r.get("statement"))
            for r in expenses
        ],
        paste_cols=6,
        widths=[12, 26, 11, 22, 7, 20, 12, 50, 40, 14, 24],
    )
    _category_dropdown(ws, "D", f"Categories!$A$1:$A${len(setup.expense_categories)}", len(expenses) + 1)

    ws = wb.create_sheet("Income", 1)
    _write_sheet(
        ws,
        ["Date", "Source", "$ Amount", "Income Category", "Hide?", "Notes (Optional)",
         "Statement description", "Review note", "Categorized by", "Statement"],
        [
            (_as_date(r["date"]), r["source"], r["amount"], r["category"], None, None,
             r["statement_description"], r["review_note"], r.get("categorized_by"), r.get("statement"))
            for r in income
        ],
        paste_cols=6,
        widths=[12, 30, 11, 22, 7, 20, 50, 40, 14, 24],
    )
    _category_dropdown(ws, "D", f"Categories!$B$1:$B${len(setup.income_categories)}", len(income) + 1)

    ws = wb.create_sheet("Skipped", 2)
    _write_sheet(
        ws,
        ["Date", "Statement description", "Amount", "Reason", "Categorized by", "Statement"],
        [
            (_as_date(r["date"]), r["statement_description"], r["amount"],
             r["reason"], r.get("categorized_by"), r.get("statement"))
            for r in skipped
        ],
        paste_cols=0,
        widths=[12, 50, 11, 40, 14, 24],
    )

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
