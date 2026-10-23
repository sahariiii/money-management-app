"""Read categories and past entries from a 'Budget tracking tool' workbook.

Category names are read from the 'Category Setup' tab, so the agent always
uses the user's current list, and past rows on 'Expenses' / 'Income' serve as
examples of how they name vendors and pick categories.

An optional 'Vendor Rules' tab maps statement text to a fixed vendor and
category; those rows are categorized without the agent. Its first row holds
the headers Match, Vendor, Category, Type (Expense / Income / Skip), Shared?
and Note, in any order.
"""

import io
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Literal

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

# How many recent past entries per sheet to show the agent as examples.
MAX_EXAMPLES = 200


@dataclass(frozen=True)
class Category:
    name: str
    note: str | None


@dataclass(frozen=True)
class PastEntry:
    date: date
    vendor: str
    amount: float
    category: str
    note: str | None


@dataclass(frozen=True)
class VendorRule:
    match: str  # case-insensitive text to find in the statement description
    vendor: str | None  # vendor (expenses) or source (income); unused for skip rules
    category: str | None  # unused for skip rules
    kind: Literal["expense", "income", "skip"]
    shared: bool
    note: str | None  # the skip reason for skip rules


@dataclass(frozen=True)
class BudgetSetup:
    expense_categories: list[Category]
    income_categories: list[Category]
    past_expenses: list[PastEntry]
    past_income: list[PastEntry]
    rules: list[VendorRule] = field(default_factory=list)


class RulesError(ValueError):
    pass


def _categories(ws: Worksheet, header_prefix: str) -> list[Category]:
    """Category names listed under a column-B header such as 'Expense Categories (up to 80)'."""
    rows = list(ws.iter_rows(min_col=2, max_col=4))
    start = next(i for i, row in enumerate(rows) if str(row[0].value or "").startswith(header_prefix))
    categories: list[Category] = []
    for name_cell, _, note_cell in rows[start + 1 :]:
        name = name_cell.value
        if isinstance(name, str) and name.startswith("="):
            break  # reached the formula-driven summary tables further down
        if isinstance(name, str) and "categories" in name.lower():
            break  # reached the next section's header
        if name not in (None, ""):
            note = note_cell.value if note_cell.value not in (None, "", name) else None
            categories.append(Category(str(name).strip(), note))
    return categories


def _past_entries(ws: Worksheet) -> list[PastEntry]:
    entries: list[PastEntry] = []
    # Data starts under the header row at row 7; columns B-G are date, vendor, amount, category, hide, notes.
    for row in ws.iter_rows(min_row=8, min_col=2, max_col=7, values_only=True):
        when, vendor, amount, category, hide, note = row
        if not isinstance(when, datetime) or vendor is None or category is None or hide:
            continue
        try:
            amount = float(amount)
        except (TypeError, ValueError):
            continue
        entries.append(PastEntry(when.date(), str(vendor).strip(), amount, str(category), note or None))
    entries.sort(key=lambda e: e.date, reverse=True)
    return entries[:MAX_EXAMPLES]


def _text(value) -> str | None:
    text = str(value).strip() if value is not None else ""
    return text or None


def _vendor_rules(ws: Worksheet, expense: list[Category], income: list[Category]) -> list[VendorRule]:
    rows = ws.iter_rows(values_only=True)
    headers = [(_text(h) or "").lower().rstrip("?") for h in next(rows, ())]
    columns = {name: headers.index(name) for name in ("match", "vendor", "category", "type", "shared", "note") if name in headers}
    if "match" not in columns or "type" not in columns:
        raise RulesError("The 'Vendor Rules' tab needs at least 'Match' and 'Type' headers in its first row.")

    valid = {"expense": {c.name for c in expense}, "income": {c.name for c in income}}
    rules: list[VendorRule] = []
    problems: list[str] = []
    for row_number, row in enumerate(rows, start=2):
        def cell(name: str) -> str | None:
            i = columns.get(name)
            return _text(row[i]) if i is not None and i < len(row) else None

        match = cell("match")
        if not match:
            continue
        kind = (cell("type") or "").lower()
        if kind not in ("expense", "income", "skip"):
            problems.append(f"row {row_number}: Type must be Expense, Income or Skip, not {cell('type')!r}")
            continue
        vendor, category = cell("vendor"), cell("category")
        if kind != "skip":
            if not vendor:
                problems.append(f"row {row_number}: {kind} rule '{match}' has no Vendor")
            if category not in valid[kind]:
                problems.append(f"row {row_number}: {category!r} is not one of your {kind} categories")
            if not vendor or category not in valid[kind]:
                continue
        shared = (cell("shared") or "").lower() in ("yes", "y", "true", "x", "1")
        rules.append(VendorRule(match, vendor, category, kind, shared, cell("note")))

    if problems:
        raise RulesError("Fix these rows on the 'Vendor Rules' tab: " + "; ".join(problems))
    return rules


def load_budget_setup(source: bytes | str | Path) -> BudgetSetup:
    data = io.BytesIO(source) if isinstance(source, bytes) else source
    # data_only would return cached formula results; categories are plain values either way.
    # read_only streams just the sheets used below instead of parsing every sheet up front.
    wb = openpyxl.load_workbook(data, read_only=True)
    try:
        setup = wb["Category Setup"]
        expense = _categories(setup, "Expense Categories")
        income = _categories(setup, "Income Categories")
        return BudgetSetup(
            expense_categories=expense,
            income_categories=income,
            past_expenses=_past_entries(wb["Expenses"]),
            past_income=_past_entries(wb["Income"]),
            rules=_vendor_rules(wb["Vendor Rules"], expense, income) if "Vendor Rules" in wb.sheetnames else [],
        )
    finally:
        wb.close()
